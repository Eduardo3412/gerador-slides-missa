from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

import re


# DIVIDIR CONTEÚDO
def dividir_blocos_automaticamente(texto):

    blocos_originais = re.split(
        r"\r?\n\s*\r?\n",
        texto.strip()
    )

    slides = []

    for bloco in blocos_originais:

        linhas = [
            linha.strip()
            for linha in bloco.splitlines()
            if linha.strip()
        ]

        # Máximo de 4 linhas por slide
        for i in range(0, len(linhas), 4):

            grupo = linhas[i:i + 4]

            slides.append(
                "\n".join(grupo)
            )

    return slides


# CRIAR SLIDE DE TEXTO
def criar_slide_texto(
    apresentacao,
    texto
):

    slide = apresentacao.slides.add_slide(
        apresentacao.slide_layouts[6]
    )

    # Fundo preto
    fundo = slide.background

    preenchimento = fundo.fill
    preenchimento.solid()

    preenchimento.fore_color.rgb = RGBColor(
        0,
        0,
        0
    )


    # Caixa de texto
    caixa = slide.shapes.add_textbox(
        Inches(0.5),
        Inches(0.5),
        Inches(12.333),
        Inches(6.5)
    )

    quadro = caixa.text_frame

    quadro.clear()
    quadro.word_wrap = True

    quadro.vertical_anchor = MSO_ANCHOR.MIDDLE


    paragrafo = quadro.paragraphs[0]

    paragrafo.text = texto

    paragrafo.alignment = PP_ALIGN.CENTER


    for run in paragrafo.runs:

        run.font.name = "Arial"

        run.font.size = Pt(60)

        run.font.bold = True

        run.font.color.rgb = RGBColor(
            255,
            255,
            255
        )


# CRIAR SLIDE PRETO
def criar_slide_preto(apresentacao):

    slide = apresentacao.slides.add_slide(
        apresentacao.slide_layouts[6]
    )

    fundo = slide.background

    preenchimento = fundo.fill
    preenchimento.solid()

    preenchimento.fore_color.rgb = RGBColor(
        0,
        0,
        0
    )


# GERAR APRESENTAÇÃO COMPLETA
def gerar_apresentacao_completa(
    momentos,
    adicionar_slide_preto,
    nome_arquivo
):

    apresentacao = Presentation()


    # Formato 16:9
    apresentacao.slide_width = Inches(13.333)

    apresentacao.slide_height = Inches(7.5)


    # Percorre todos os momentos
    for momento in momentos:

        conteudo = momento.conteudo


        # Divide o conteúdo
        blocos = dividir_blocos_automaticamente(
            conteudo
        )


        # Cria os slides
        for bloco in blocos:

            criar_slide_texto(
                apresentacao,
                bloco
            )


        # Slide preto ao final do momento
        if adicionar_slide_preto:

            criar_slide_preto(
                apresentacao
            )


    # Salva o arquivo
    apresentacao.save(
        nome_arquivo
    ) 
